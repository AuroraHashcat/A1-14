"""文档导出工具"""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.models import Report

INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`|~~[^~]+~~)")


def _replace_markdown_links(text: str) -> str:
    """Convert markdown links [text](url) to 'text (url)'."""

    def _repl(match: re.Match[str]) -> str:
        label = match.group(1).strip()
        url = match.group(2).strip()
        return f"{label} ({url})" if url else label

    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", _repl, text)


def _append_runs(paragraph, text: str) -> None:
    """Add runs to paragraph while handling basic inline markdown tokens."""
    if not text:
        return

    normalized = _replace_markdown_links(text)
    parts = INLINE_PATTERN.split(normalized)

    for part in parts:
        if not part:
            continue

        run_text = part
        run = paragraph.add_run()

        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run_text = part[2:-2]
            run.bold = True
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            run_text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_text = part[1:-1]
            run.italic = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            run_text = part[1:-1]
            run.italic = True
        elif part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            run_text = part[2:-2]
            run.font.strike = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run_text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(10)

        run.text = run_text


def _add_paragraph_with_style(document: Document, text: str, style: str | None = None, indent_level: int = 0):
    paragraph = document.add_paragraph(style=style)
    if indent_level > 0:
        paragraph.paragraph_format.left_indent = Pt(18 * indent_level)
    _append_runs(paragraph, text)
    return paragraph


def _add_code_block(document: Document, code: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def _configure_document(document: Document) -> None:
    """设置文档的基础字体，确保中文显示正常。"""
    style = document.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")  # type: ignore[attr-defined]


def _add_markdownish_content(document: Document, content: str) -> None:
    """粗略地将 Markdown 文本转为 docx 段落。"""
    if not content:
        return

    lines = content.replace("\r\n", "\n").split("\n")
    in_code_block = False
    code_buffer: list[str] = []

    for raw_line in lines:
        stripped_line = raw_line.rstrip("\n")
        if stripped_line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_buffer = []
            else:
                _add_code_block(document, "\n".join(code_buffer))
                code_buffer = []
                in_code_block = False
            continue

        if in_code_block:
            code_buffer.append(stripped_line)
            continue

        line = stripped_line.rstrip()
        if not line:
            document.add_paragraph("")
            continue

        stripped = line.lstrip()
        indent_level = (len(line) - len(stripped)) // 2

        if stripped.startswith("- ") or stripped.startswith("* "):
            text = stripped[2:].strip()
            _add_paragraph_with_style(document, text, style="List Bullet", indent_level=indent_level)
            continue

        ordered_match = re.match(r"^(\d+)[\.)]\s+(.*)", stripped)
        if ordered_match:
            _add_paragraph_with_style(
                document,
                ordered_match.group(2).strip(),
                style="List Number",
                indent_level=indent_level,
            )
            continue

        if stripped.startswith(">"):
            quote_text = stripped.lstrip(">").strip()
            paragraph = _add_paragraph_with_style(document, quote_text)
            try:
                paragraph.style = document.styles["Intense Quote"]
            except KeyError:
                pass
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_text = stripped[level:].strip() or stripped.strip("# ")
            heading_paragraph = document.add_heading("", level=min(level, 4))
            _append_runs(heading_paragraph, heading_text)
            continue

        _add_paragraph_with_style(document, stripped)

    if in_code_block and code_buffer:
        _add_code_block(document, "\n".join(code_buffer))


def _sanitize_filename(name: str, suffix: str) -> str:
    base = re.sub(r"[\s]+", "_", name.strip()) if name else "export"
    base = re.sub(r"[^\w\-\u4e00-\u9fff]", "", base)
    base = base or "export"
    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    return f"{base}_{timestamp}.{suffix}"


def build_report_docx(report: Report) -> Tuple[BytesIO, str]:
    """基于报告内容生成 Word 文档。"""
    document = Document()
    _configure_document(document)

    title = report.title or "密评自动化报告"
    document.add_heading(title, level=0)

    system_name = report.evaluation.system.name if report.evaluation and report.evaluation.system else ""
    document.add_paragraph(f"系统名称：{system_name}")
    document.add_paragraph(f"报告编号：{report.uuid}")
    document.add_paragraph(
        f"生成时间：{(report.created_at or datetime.utcnow()).strftime('%Y-%m-%d %H:%M:%S')}"
    )
    if report.score is not None:
        document.add_paragraph(f"评测得分：{round(report.score, 2)}")

    if report.summary:
        document.add_heading("报告摘要", level=1)
        _add_markdownish_content(document, report.summary)

    document.add_heading("报告正文", level=1)
    _add_markdownish_content(document, report.content or "")

    if report.recommendations:
        document.add_heading("改进建议", level=1)
        if isinstance(report.recommendations, list):
            recs = report.recommendations
        elif isinstance(report.recommendations, str):
            recs = [report.recommendations]
        else:
            recs = []
        for item in recs:
            if isinstance(item, str):
                _add_paragraph_with_style(document, item, style="List Bullet")
            elif isinstance(item, dict):
                _add_paragraph_with_style(
                    document,
                    item.get("description") or item.get("content") or str(item),
                    style="List Bullet",
                )

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    filename = _sanitize_filename(title, "docx")
    return buffer, filename


def build_question_bank_docx(
    topic: str,
    difficulty: str,
    questions: Iterable[Dict[str, Any]],
) -> Tuple[BytesIO, str]:
    """将生成的题库导出为 Word 文档。"""
    document = Document()
    _configure_document(document)

    title = f"密评题库 - {topic}" if topic else "密评题库"
    document.add_heading(title, level=0)

    difficulty_map = {
        "easy": "基础",
        "medium": "进阶",
        "hard": "挑战",
    }
    document.add_paragraph(f"难度：{difficulty_map.get(difficulty, difficulty or '未知')}")
    document.add_paragraph(f"导出时间：{datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC")

    for index, question in enumerate(questions, start=1):
        document.add_heading(f"题目 {index}", level=1)
        _add_paragraph_with_style(document, question.get("question") or "(题干缺失)")

        options = question.get("options") or question.get("choices")
        if isinstance(options, dict):
            for key, value in options.items():
                _add_paragraph_with_style(document, f"{key}. {value}", style="List Bullet")
        elif isinstance(options, list):
            for option_index, value in enumerate(options, start=1):
                label = chr(64 + option_index) if option_index <= 26 else str(option_index)
                _add_paragraph_with_style(document, f"{label}. {value}", style="List Bullet")

        answer = question.get("answer") or question.get("correct_answer")
        if answer:
            _add_paragraph_with_style(document, f"参考答案：{answer}")

        explanation = question.get("explanation") or question.get("analysis")
        if explanation:
            _add_paragraph_with_style(document, f"答案解析：{explanation}")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    filename = _sanitize_filename(title, "docx")
    return buffer, filename