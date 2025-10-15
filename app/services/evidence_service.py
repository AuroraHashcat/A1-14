"""
证据处理服务
"""

import os
import fitz  # PyMuPDF
from PIL import Image
import easyocr
from docx import Document
from app.models import Evidence, db
from typing import Optional, Dict, Any
import json
import asyncio
from concurrent.futures import ThreadPoolExecutor

class EvidenceService:
    """证据文件处理服务"""
    
    def __init__(self):
        self.ocr_reader = None
        self.executor = ThreadPoolExecutor(max_workers=2)
    
    def _get_ocr_reader(self):
        """延迟初始化OCR读取器"""
        if self.ocr_reader is None:
            self.ocr_reader = easyocr.Reader(['ch_sim', 'en'])
        return self.ocr_reader
    
    def extract_content(self, evidence_id: int) -> Dict[str, Any]:
        """提取证据文件内容"""
        evidence = Evidence.query.get(evidence_id)
        if not evidence:
            raise ValueError(f"Evidence with id {evidence_id} not found")
        
        try:
            # 根据文件类型选择处理方法
            if evidence.file_type.startswith('image/'):
                extracted_text = self._extract_from_image(evidence.file_path)
                metadata = self._get_image_metadata(evidence.file_path)
            elif evidence.file_type == 'application/pdf':
                extracted_text = self._extract_from_pdf(evidence.file_path)
                metadata = self._get_pdf_metadata(evidence.file_path)
            elif evidence.file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                extracted_text = self._extract_from_docx(evidence.file_path)
                metadata = self._get_docx_metadata(evidence.file_path)
            else:
                extracted_text = ""
                metadata = {}
            
            # 更新数据库
            evidence.extracted_text = extracted_text
            evidence.metadata = metadata
            db.session.commit()
            
            return {
                'evidence_id': evidence_id,
                'extracted_text': extracted_text,
                'metadata': metadata,
                'success': True
            }
            
        except Exception as e:
            return {
                'evidence_id': evidence_id,
                'error': str(e),
                'success': False
            }
    
    def extract_content_async(self, evidence_id: int):
        """异步提取证据文件内容"""
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        future = self.executor.submit(self.extract_content, evidence_id)
        return future
    
    def _extract_from_image(self, file_path: str) -> str:
        """从图片中提取文本"""
        try:
            ocr_reader = self._get_ocr_reader()
            results = ocr_reader.readtext(file_path)
            
            # 合并OCR结果
            text_lines = []
            for (bbox, text, confidence) in results:
                if confidence > 0.5:  # 置信度阈值
                    text_lines.append(text)
            
            return '\n'.join(text_lines)
            
        except Exception as e:
            print(f"OCR extraction error: {e}")
            return ""
    
    def _extract_from_pdf(self, file_path: str) -> str:
        """从PDF中提取文本"""
        try:
            doc = fitz.open(file_path)
            text_content = []
            
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()
                text_content.append(text)
            
            doc.close()
            return '\n'.join(text_content)
            
        except Exception as e:
            print(f"PDF extraction error: {e}")
            return ""
    
    def _extract_from_docx(self, file_path: str) -> str:
        """从Word文档中提取文本"""
        try:
            doc = Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                text_content.append(paragraph.text)
            
            return '\n'.join(text_content)
            
        except Exception as e:
            print(f"DOCX extraction error: {e}")
            return ""
    
    def _get_image_metadata(self, file_path: str) -> Dict[str, Any]:
        """获取图片元数据"""
        try:
            with Image.open(file_path) as img:
                return {
                    'format': img.format,
                    'mode': img.mode,
                    'size': img.size,
                    'has_exif': bool(getattr(img, '_getexif', None))
                }
        except:
            return {}
    
    def _get_pdf_metadata(self, file_path: str) -> Dict[str, Any]:
        """获取PDF元数据"""
        try:
            doc = fitz.open(file_path)
            metadata = doc.metadata
            page_count = len(doc)
            doc.close()
            
            return {
                'page_count': page_count,
                'title': metadata.get('title', ''),
                'author': metadata.get('author', ''),
                'creator': metadata.get('creator', ''),
                'producer': metadata.get('producer', ''),
                'creation_date': metadata.get('creationDate', ''),
                'modification_date': metadata.get('modDate', '')
            }
        except:
            return {}
    
    def _get_docx_metadata(self, file_path: str) -> Dict[str, Any]:
        """获取Word文档元数据"""
        try:
            doc = Document(file_path)
            core_props = doc.core_properties
            
            return {
                'author': core_props.author or '',
                'category': core_props.category or '',
                'comments': core_props.comments or '',
                'content_status': core_props.content_status or '',
                'created': core_props.created.isoformat() if core_props.created else '',
                'identifier': core_props.identifier or '',
                'keywords': core_props.keywords or '',
                'language': core_props.language or '',
                'last_modified_by': core_props.last_modified_by or '',
                'last_printed': core_props.last_printed.isoformat() if core_props.last_printed else '',
                'modified': core_props.modified.isoformat() if core_props.modified else '',
                'revision': core_props.revision,
                'subject': core_props.subject or '',
                'title': core_props.title or '',
                'version': core_props.version or ''
            }
        except:
            return {}
